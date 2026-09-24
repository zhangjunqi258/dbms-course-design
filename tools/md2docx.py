# -*- coding: utf-8 -*-
"""把课程设计报告的 Markdown 转换成 Word（.docx）。

支持：# ## ### 标题、段落、- 无序列表、1. 有序列表、> 引用、
| 表格 |、``` 代码块、**加粗**、--- 分页线。

用法：
    python tools/md2docx.py docs/课程设计报告.md docs/课程设计报告.docx
"""
import re
import sys
import os

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

CN_FONT = "宋体"
CN_HEAD_FONT = "黑体"
EN_FONT = "Times New Roman"
CODE_FONT = "Consolas"


def set_run_font(run, cn=CN_FONT, en=EN_FONT, size=None, bold=None):
    run.font.name = en
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def init_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = EN_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.4
    pf.space_after = Pt(4)

    for name, size in (("Heading 1", 16), ("Heading 2", 14),
                       ("Heading 3", 12), ("Title", 22)):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = EN_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_HEAD_FONT)


BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def add_rich_text(para, text, cn=CN_FONT, en=EN_FONT, size=None):
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = para.add_run(text[pos:m.start()])
            set_run_font(r, cn, en, size)
        r = para.add_run(m.group(1))
        set_run_font(r, cn, en, size, bold=True)
        pos = m.end()
    if pos < len(text):
        r = para.add_run(text[pos:])
        set_run_font(r, cn, en, size)
    if pos == 0 and not BOLD_RE.search(text):
        pass


def split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def is_sep_row(line):
    s = line.strip().strip("|")
    return bool(s) and all(ch in "-: " for ch in s)


def convert(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    init_styles(doc)

    i = 0
    in_code = False
    code_buf = []
    in_cover = True
    n = len(lines)

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 代码块
        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run("\n".join(code_buf))
                set_run_font(r, CODE_FONT, CODE_FONT, 9)
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # 分页线
        if stripped == "---":
            if in_cover:
                doc.add_page_break()
                in_cover = False
            else:
                p = doc.add_paragraph()
                set_run_font(p.add_run("─" * 30), CN_FONT, EN_FONT, 9)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                p = doc.add_heading(level=0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_run_font(p.add_run(text), CN_HEAD_FONT, EN_FONT, 22, bold=True)
            else:
                p = doc.add_heading(level=min(level - 1, 3))
                set_run_font(p.add_run(text), CN_HEAD_FONT, EN_FONT,
                             16 if level == 2 else 14 if level == 3 else 12,
                             bold=True)
            i += 1
            continue

        # 表格
        if stripped.startswith("|"):
            block = []
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [split_row(b) for b in block if not is_sep_row(b)]
            if rows:
                cols = max(len(r) for r in rows)
                t = doc.add_table(rows=0, cols=cols)
                t.style = "Table Grid"
                for ri, row in enumerate(rows):
                    cells = t.add_row().cells
                    for ci in range(cols):
                        txt = row[ci] if ci < len(row) else ""
                        para = cells[ci].paragraphs[0]
                        para.paragraph_format.line_spacing = 1.0
                        set_run_font(para.add_run(txt), CN_FONT, EN_FONT, 9,
                                     bold=(ri == 0))
                doc.add_paragraph()
            continue

        # 引用
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            add_rich_text(p, stripped.lstrip("> ").strip(), CN_FONT, EN_FONT, 10)
            i += 1
            continue

        # 列表
        m = re.match(r"^([-*])\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, m.group(2))
            i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_rich_text(p, m.group(2))
            i += 1
            continue

        # 封面：居中的普通段落
        if in_cover and not stripped.startswith("**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_rich_text(p, stripped, CN_FONT, EN_FONT, 12)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_rich_text(p, stripped)
        i += 1

    doc.save(docx_path)
    print("已生成：", docx_path)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__)
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])
